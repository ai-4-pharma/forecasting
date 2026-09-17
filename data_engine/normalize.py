"""data_engine.normalize — Leitura, mapeamento e normalização canônica (longo/largo).
(extraído de data_engine.py single-file; split 16/09/2026).
"""

from __future__ import annotations

from datetime import date
import io
import json
import polars as pl
import polars.selectors as cs
import openpyxl
from contracts import (
    CanonicalDataset,
    Layout,
    MAX_CANONICAL_OBSERVATIONS,
    MAX_ENTITIES,
    MappingConfig,
    Severity,
    StudyConfig,
    ValidationIssue,
    empty_long,
    stable_id,
)
from .dates import parse_date_main
from .ingest import _detect_delimiter, _sniff_encoding


def _read_dataframe(
    content: bytes, filename: str, mapping: MappingConfig, issues: list[ValidationIssue]
) -> pl.DataFrame:
    name = filename.lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = mapping.sheet_name or wb.sheetnames[0]
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            issues.append(
                ValidationIssue(
                    code="E_ROWS_INVALID",
                    severity=Severity.ERROR,
                    message="Planilha vazia.",
                )
            )
            return empty_long("col", 1)
        col_names = [("" if v is None else str(v)) for v in rows[0]]
        data_rows = rows[1:]
        return pl.DataFrame(
            data_rows, schema={c: pl.Utf8 for c in col_names}, orient="row"
        ).with_columns(cs.all().cast(pl.Utf8).fill_null(""))

    enc = mapping.encoding or _sniff_encoding(content)
    raw = content.decode(enc, errors="replace")
    del_ = mapping.delimiter or _detect_delimiter(raw)
    try:
        df = pl.read_csv(
            io.StringIO(raw),
            separator=del_,
            has_header=True,
            infer_schema_length=0,
            truncate_ragged_lines=True,
            ignore_errors=False,
        )
    except Exception as e:  # noqa: BLE001
        issues.append(
            ValidationIssue(
                code="E_ROWS_INVALID",
                severity=Severity.ERROR,
                message=f"Falha ao ler CSV: {e}",
            )
        )
        return empty_long("col", 1)
    # força string em todas as colunas para decisões de parser seguras
    return df.with_columns(cs.all().cast(pl.Utf8).fill_null(""))


def _parse_number(text: str, mapping: MappingConfig) -> float | None:
    s = str(text).strip()
    if s in ("", "-", "--", "n/a", "N/A"):
        return None
    if mapping.thousands_separator and mapping.thousands_separator in s:
        s = s.replace(mapping.thousands_separator, "")
    s = s.replace("R$", "").replace(" ", "").replace("\\xa0", "")
    if mapping.decimal_separator == ",":
        if "." in s and "," in s:
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def entity_matrix_stable(mapping: MappingConfig, row: dict, key_cols: list[str]) -> str:
    """Hash SHA-256 estavel da chave (doc 5.2)."""
    pairs = [(c, str(row.get(c, "") or "")) for c in mapping.key_json_order or key_cols]
    return stable_id(*pairs)


def normalize_file(
    content: bytes, filename: str, study: StudyConfig, mapping: MappingConfig
) -> CanonicalDataset:
    """Normaliza CSV/XLSX para observacoes canonicas (sec 7.3)."""
    issues: list[ValidationIssue] = []
    df = _read_dataframe(content, filename, mapping, issues)
    if df.height == 0 or any(
        i.severity == Severity.ERROR for i in issues if i.code == "E_ROWS_INVALID"
    ):
        return CanonicalDataset(
            empty_long("source_row_id", 0),
            empty_long("entity_id", 0),
            empty_long("source_row_id", 0),
            issues,
            stable_id("empty"),
        )

    # conferir colunas essenciais
    missing_key = [c for c in mapping.key_columns if c not in df.columns]
    if missing_key:
        issues.append(
            ValidationIssue(
                code="E_KEY_MISSING",
                severity=Severity.ERROR,
                message=f"Colunas da chave ausentes: {missing_key}",
            )
        )
    if study.layout == Layout.LONG:
        if not mapping.period_column or mapping.period_column not in df.columns:
            issues.append(
                ValidationIssue(
                    code="E_PERIOD_COLUMN_MISSING",
                    severity=Severity.ERROR,
                    message="Coluna de período não mapeada.",
                )
            )
        for measure, col in mapping.measure_columns.items():
            if col not in df.columns:
                issues.append(
                    ValidationIssue(
                        code="E_MEASURE_MISSING",
                        severity=Severity.ERROR,
                        message=f"Coluna de medida '{col}' ausente.",
                    )
                )
    elif not mapping.wide_period_map:
        issues.append(
            ValidationIssue(
                code="E_WIDE_PERIOD_INCOMPLETE",
                severity=Severity.ERROR,
                message="Nenhum período mapeado para o formato largo.",
            )
        )
    if any(i.severity == Severity.ERROR for i in issues):
        return CanonicalDataset(
            empty_long("source_row_id", 0),
            empty_long("entity_id", 0),
            empty_long("source_row_id", 0),
            issues,
            stable_id("nok"),
        )

    rows = df.to_dicts()
    observed_rows: list[dict] = []
    entities: dict[str, dict] = {}
    obs_rows: list[dict] = []
    source_rows: list[dict] = []
    entity_keys: dict[str, tuple] = {}
    dim_pool: dict[str, set] = {}

    if study.layout == Layout.LONG:
        period_col = mapping.period_column
        for i, row in enumerate(rows):
            src_id = f"r{i}"
            source_rows.append(
                {
                    "dataset_id": "",  # preenchido no save
                    "source_row_id": src_id,
                    "sheet": mapping.sheet_name or filename,
                    "row_number": i + 2,
                    "raw_json": json.dumps(
                        {k: v for k, v in row.items()}, ensure_ascii=False
                    ),
                }
            )
            period_col = mapping.period_column or ""
            period_text = row.get(period_col, "")
            ds = parse_date_main(
                period_text, mapping.date_format, study.source_frequency
            )
            key_vals = tuple(row.get(c, "") or "" for c in mapping.key_columns)
            if not any(key_vals):
                issues.append(
                    ValidationIssue(
                        code="E_KEY_EMPTY",
                        severity=Severity.ERROR,
                        message="Linha sem valor na chave.",
                        source_row_ids=[src_id],
                    )
                )
                continue
            if ds is None:
                issues.append(
                    ValidationIssue(
                        code="E_DATE_AMBIGUOUS",
                        severity=Severity.ERROR,
                        message=f"Data não interpretada: '{period_text}'",
                        source_row_ids=[src_id],
                    )
                )
                continue
            entity_id = entity_matrix_stable(mapping, row, mapping.key_columns)
            dims: dict[str, str] = {}
            for logical, col in mapping.dimension_columns.items():
                dims[logical] = row.get(col, "") or ""
            attrs: dict[str, str] = {}
            for logical, col in mapping.attribute_columns.items():
                attrs[logical] = row.get(col, "") or ""
            _store_entity(
                entities,
                entity_keys,
                dim_pool,
                entity_id,
                key_vals,
                dims,
                attrs,
                src_id,
                issues,
                mapping,
                study,
            )
            for measure, col in mapping.measure_columns.items():
                num = _parse_number(row.get(col, ""), mapping)
                if num is None:
                    issues.append(
                        ValidationIssue(
                            code="E_ROW_INVALID",
                            severity=Severity.WARNING,
                            message="Valor de medida não numérico.",
                            source_row_ids=[src_id],
                        )
                    )
                    continue
                if measure == "unidades" and not float(num).is_integer():
                    issues.append(
                        ValidationIssue(
                            code="E_UNITS_FRACTION",
                            severity=Severity.ERROR,
                            message="Unidades devem ser inteiras.",
                            source_row_ids=[src_id],
                        )
                    )
                    continue
                series_id = stable_id(entity_id, measure)
                obs_rows.append(
                    {
                        "source_row_id": src_id,
                        "entity_id": entity_id,
                        "series_id": series_id,
                        "ds": ds,
                        "measure": measure,
                        "y_raw": num,
                        "observed": True,
                    }
                )
    else:
        # formato largo
        period_map = mapping.wide_period_map  # origem -> ISO
        measures = mapping.measure_columns
        for i, row in enumerate(rows):
            src_id = f"r{i}"
            source_rows.append(
                {
                    "dataset_id": "",
                    "source_row_id": src_id,
                    "sheet": mapping.sheet_name or filename,
                    "row_number": i + 2,
                    "raw_json": json.dumps(dict(row), ensure_ascii=False),
                }
            )
            key_vals = tuple(row.get(c, "") or "" for c in mapping.key_columns)
            if not any(key_vals):
                issues.append(
                    ValidationIssue(
                        code="E_KEY_EMPTY",
                        severity=Severity.ERROR,
                        message="Linha sem chave.",
                        source_row_ids=[src_id],
                    )
                )
                continue
            entity_id = entity_matrix_stable(mapping, row, mapping.key_columns)
            dims = {
                logical: row.get(col, "") or ""
                for logical, col in mapping.dimension_columns.items()
            }
            attrs = {
                logical: row.get(col, "") or ""
                for logical, col in mapping.attribute_columns.items()
            }
            _store_entity(
                entities,
                entity_keys,
                dim_pool,
                entity_id,
                key_vals,
                dims,
                attrs,
                src_id,
                issues,
                mapping,
                study,
            )

            # medida da linha: rótulo ou medida única (sem coluna de medida)
            line_measure: str | None = None
            if (
                mapping.measure_label_column
                and mapping.measure_label_column in df.columns
            ):
                raw_m = (row.get(mapping.measure_label_column) or "").strip().lower()
                valid = {
                    "unidades": "unidades",
                    "unidades": "unidades",
                    "un": "unidades",
                    "unidade": "unidades",
                    "valor": "valor",
                    "venda": "valor",
                    "value": "valor",
                    "val": "valor",
                }
                line_measure = valid.get(raw_m, None)
                if line_measure is None:
                    issues.append(
                        ValidationIssue(
                            code="E_MEASURE_UNKNOWN",
                            severity=Severity.ERROR,
                            message=f"Rótulo de medida indefinido: {raw_m}",
                            source_row_ids=[src_id],
                        )
                    )
                    continue
            elif len(mapping.measure_columns) == 1:
                line_measure = next(iter(mapping.measure_columns))
            else:
                issues.append(
                    ValidationIssue(
                        code="E_MEASURE_UNKNOWN",
                        severity=Severity.ERROR,
                        message="Formato largo com 2 medidas exige "
                        "coluna de rótulo da medida.",
                        source_row_ids=[src_id],
                    )
                )
                continue

            for src_col, ds_iso in period_map.items():
                ds = date.fromisoformat(ds_iso)
                raw_v = row.get(src_col, "")
                num = _parse_number(raw_v, mapping) if raw_v not in (None, "") else None
                if num is None:
                    issues.append(
                        ValidationIssue(
                            code="E_ROW_INVALID",
                            severity=Severity.WARNING,
                            message="Valor não numérico no período.",
                            source_row_ids=[src_id],
                        )
                    )
                    continue
                if line_measure == "unidades" and not float(num).is_integer():
                    issues.append(
                        ValidationIssue(
                            code="E_UNITS_FRACTION",
                            severity=Severity.ERROR,
                            message="Unidades devem ser inteiras.",
                            source_row_ids=[src_id],
                        )
                    )
                    continue
                series_id = stable_id(entity_id, line_measure)
                obs_rows.append(
                    {
                        "source_row_id": src_id,
                        "entity_id": entity_id,
                        "series_id": series_id,
                        "ds": ds,
                        "measure": line_measure,
                        "y_raw": float(num),
                        "observed": True,
                    }
                )

    if len(entities) > MAX_ENTITIES:
        issues.append(
            ValidationIssue(
                code="E_ENTITY_LIMIT",
                severity=Severity.ERROR,
                message=f"{len(entities)} entidades, acima do limite "
                f"de {MAX_ENTITIES}.",
            )
        )
    if len(obs_rows) > MAX_CANONICAL_OBSERVATIONS:
        issues.append(
            ValidationIssue(
                code="E_OBS_LIMIT",
                severity=Severity.ERROR,
                message=f"{len(obs_rows)} observações, acima do limite "
                f"de {MAX_CANONICAL_OBSERVATIONS}.",
            )
        )

    obs_df = pl.DataFrame(
        obs_rows
        or [
            {
                "source_row_id": "x",
                "entity_id": "x",
                "series_id": "x",
                "ds": date(2020, 1, 1),
                "measure": "unidades",
                "y_raw": 0.0,
                "observed": True,
            }
        ]
    )
    obs_df = (
        obs_df.with_columns(
            pl.col("y_raw")
            .cast(pl.Float64)
            .map_elements(lambda v: round(v, 2), return_dtype=pl.Float64)
            .alias("y_raw_dec")
        )
        .drop("y_raw")
        .rename({"y_raw_dec": "y_raw"})
        .with_columns(pl.col("y_raw").cast(pl.Decimal(20, 2)))
    )
    ent_df = pl.DataFrame(
        [
            {
                "entity_id": eid,
                "key_json": json.dumps(dict(zip(mapping.key_columns, keys))),
                "dimensions_json": json.dumps(dims, ensure_ascii=False),
                "attributes_json": json.dumps(attrs, ensure_ascii=False),
            }
            for eid, (keys, dims, attrs) in entities.items()
        ]
    )
    src_df = pl.DataFrame(
        [
            {k: (v if not isinstance(v, dict) else json.dumps(v)) for k, v in r.items()}
            for r in source_rows
        ]
    )
    content_digest = _observations_content_digest(obs_df)
    fingerprint = stable_id(
        filename,
        str(len(obs_rows)),
        str(len(entities)),
        mapping.to_json(),
        study.to_json(),
        content_digest,
    )
    return CanonicalDataset(obs_df, ent_df, src_df, issues, fingerprint)


def _observations_content_digest(obs_df: pl.DataFrame) -> str:
    """Resumo do conteúdo real das observações (B19): o `fingerprint` do
    dataset não pode depender só da forma (nome/contagens/mapeamento), senão
    dois arquivos homônimos com dados diferentes colidem em `preparation_id`.
    """
    if obs_df.height == 0:
        return stable_id("empty_obs")
    ordered = obs_df.select("series_id", "ds", "y_raw").sort(["series_id", "ds"])
    values = [f"{r['series_id']}|{r['ds']}|{r['y_raw']}" for r in ordered.to_dicts()]
    return stable_id(*values)


def _store_entity(
    entities: dict,
    entity_keys: dict,
    dim_pool: dict,
    entity_id: str,
    key_vals: tuple,
    dims: dict,
    attrs: dict,
    src_id: str,
    issues: list[ValidationIssue],
    mapping: MappingConfig,
    study: StudyConfig,
) -> None:
    if entity_id in entities:
        prev_dims = entities[entity_id][1]
        for k, v in dims.items():
            if prev_dims.get(k, "") != v and prev_dims.get(k, ""):
                issues.append(
                    ValidationIssue(
                        code="E_DIMENSION_CONFLICT",
                        severity=Severity.ERROR,
                        message=f"Dimensão '{k}' muda entre linhas "
                        f"'{prev_dims.get(k)}' vs '{v}'.",
                        source_row_ids=[src_id],
                        entity_ids=[entity_id],
                    )
                )
    else:
        entities[entity_id] = (key_vals, dict(dims), dict(attrs))
