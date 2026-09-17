"""data_engine.ingest — Inspeção e sugestão de leitura de arquivos (separador, encoding, datas).
(extraído de data_engine.py single-file; split 16/09/2026).
"""

from __future__ import annotations

from datetime import date
import io
import re
import polars as pl
import openpyxl
from contracts import FileInspection, MAX_FILE_BYTES, MAX_PHYSICAL_ROWS, MappingConfig, Severity, ValidationIssue, empty_long



def _detect_delimiter(sample: str) -> str:
    for cand in [";", ",", "\t"]:
        lines = [l for l in sample.splitlines() if l.strip()]
        if not lines:
            continue
        if cand in lines[0]:
            return cand
    return ";"
def _sniff_encoding(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            raw.decode(enc)
            return enc if enc != "utf-8" else "utf-8"
        except UnicodeDecodeError:
            continue
    return "utf-8"
def _suggest_parse_date(text: str) -> date | None:
    """Parser mínimo de período para `suggest_time_settings` (T3.4/B14):
    aceita ISO completo, `YYYYMM`, `YYYY-MM`/`YYYY/MM` e `YYYY`."""
    c = str(text).strip()
    m = re.match(r"^(\d{4})-(\d{2})-\d{2}([ T].*)?$", c)
    if m:
        y, mth = int(m.group(1)), int(m.group(2))
        if 1 <= mth <= 12:
            return date(y, mth, 1)
    m = re.match(r"^(\d{4})(\d{2})(\.0)?$", c)
    if m:
        y, mth = int(m.group(1)), int(m.group(2))
        if 1 <= mth <= 12:
            return date(y, mth, 1)
    m = re.match(r"^(\d{4})[-/](\d{2})$", c)
    if m:
        y, mth = int(m.group(1)), int(m.group(2))
        if 1 <= mth <= 12:
            return date(y, mth, 1)
    if re.match(r"^\d{4}$", c):
        return date(int(c), 1, 1)
    return None
def suggest_time_settings(
    sample: pl.DataFrame, period_col: str | None, period_headers: list[str]
) -> dict:
    """Sugere último período fechado, quantidade de períodos e frequência
    (T3.4/B14), inferindo pelo menor espaçamento entre períodos distintos —
    substitui o default fixo `2026-08-01` da UI. `period_headers` (formato
    Largo) já vem com datas ISO (ex.: `auto_period_map` do app); `period_col`
    (formato Longo) é lido a partir da amostra de `inspect_file`.
    """
    dates: list[date] = []
    if period_headers:
        for h in period_headers:
            d = _suggest_parse_date(h)
            if d is not None:
                dates.append(d)
    elif period_col and period_col in sample.columns:
        for v in sample[period_col].drop_nulls().unique().to_list():
            d = _suggest_parse_date(str(v))
            if d is not None:
                dates.append(d)

    dates = sorted(set(dates))
    if len(dates) < 2:
        return {
            "history_end": dates[-1] if dates else None,
            "n_periods": len(dates),
            "frequency": None,
        }
    gaps = [(dates[i + 1] - dates[i]).days for i in range(len(dates) - 1)]
    min_gap = min(gaps)
    if min_gap <= 31:
        frequency = "monthly"
    elif min_gap <= 100:
        frequency = "quarterly"
    else:
        frequency = "yearly"
    return {"history_end": dates[-1], "n_periods": len(dates), "frequency": frequency}
def inspect_file(
    content: bytes, filename: str, options: MappingConfig
) -> FileInspection:
    """Inspeciona arquivo CSV/XLSX e devolve amostra e opcoes de leitura."""
    warnings: list[ValidationIssue] = []
    if len(content) > MAX_FILE_BYTES:
        warnings.append(
            ValidationIssue(
                code="E_FILE_TOO_BIG",
                severity=Severity.ERROR,
                message="Arquivo excede o limite de 100 MiB.",
            )
        )
    name = filename.lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        return _inspect_xlsx(content, filename, options, warnings)
    if name.endswith(".xls"):
        warnings.append(
            ValidationIssue(
                code="E_XLSX_FORMULA_NOT_CACHED",
                severity=Severity.ERROR,
                message="Formato .xls antigo não é suportado; salve como .xlsx.",
            )
        )
        return FileInspection([], [], empty_long("col", 0), {}, warnings, {})
    if name.endswith(".csv") or name.endswith(".txt"):
        return _inspect_csv(content, filename, options, warnings)
    warnings.append(
        ValidationIssue(
            code="E_ROW_INVALID",
            severity=Severity.ERROR,
            message="Formato de arquivo não reconhecido. Use CSV ou XLSX.",
        )
    )
    return FileInspection([], [], empty_long("col", 0), {}, warnings, {})
def _inspect_csv(
    content: bytes,
    filename: str,
    options: MappingConfig,
    warnings: list[ValidationIssue],
) -> FileInspection:
    raw = content.decode(options.encoding or _sniff_encoding(content), errors="replace")
    del_ = options.delimiter or _detect_delimiter(raw)
    sample = "\n".join(raw.splitlines()[:100])
    try:
        df = pl.read_csv(
            io.StringIO(raw),
            separator=del_,
            n_rows=100,
            infer_schema_length=100,
            has_header=True,
            truncate_ragged_lines=True,
            ignore_errors=False,
        )
    except Exception as e:  # noqa: BLE001
        warnings.append(
            ValidationIssue(
                code="E_ROWS_INVALID",
                severity=Severity.ERROR,
                message=f"Falha ao ler CSV: {e}",
            )
        )
        return FileInspection(
            [], [], empty_long("col", 0), {}, warnings, {"delimiter": del_}
        )
    physical = len(raw.splitlines())
    if physical > MAX_PHYSICAL_ROWS:
        warnings.append(
            ValidationIssue(
                code="E_ROWS_LIMIT",
                severity=Severity.ERROR,
                message=f"CSV tem {physical} linhas físicas, acima do limite de "
                f"{MAX_PHYSICAL_ROWS}.",
            )
        )
    types = {c: str(df.schema[c]) for c in df.columns}
    return FileInspection(
        [],
        list(df.columns),
        df,
        types,
        warnings,
        {"delimiter": del_, "encoding": options.encoding, "n_lines_physical": physical},
    )
def _inspect_xlsx(
    content: bytes,
    filename: str,
    options: MappingConfig,
    warnings: list[ValidationIssue],
) -> FileInspection:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        warnings.append(
            ValidationIssue(
                code="E_ROW_INVALID",
                severity=Severity.ERROR,
                message=f"Falha ao abrir XLSX: {e}",
            )
        )
        return FileInspection([], [], empty_long("col", 0), {}, warnings, {})
    sheets = wb.sheetnames
    if options.sheet_name and options.sheet_name not in sheets:
        warnings.append(
            ValidationIssue(
                code="E_ROW_INVALID",
                severity=Severity.ERROR,
                message=f"Aba '{options.sheet_name}' não existe.",
            )
        )
    sheet = options.sheet_name or sheets[0]
    ws = wb[sheet]
    if ws.max_row is None or ws.max_column is None:
        warnings.append(
            ValidationIssue(
                code="E_ROWS_INVALID",
                severity=Severity.ERROR,
                message="Planilha vazia.",
            )
        )
        return FileInspection(sheets, [], empty_long("col", 0), {}, warnings, {})
    if ws.max_row > MAX_PHYSICAL_ROWS:
        warnings.append(
            ValidationIssue(
                code="E_ROWS_LIMIT",
                severity=Severity.ERROR,
                message=f"XLSX tem {ws.max_row} linhas na aba '{sheet}', acima do limite.",
            )
        )
    rows: list[list] = []
    for r in ws.iter_rows(
        min_row=1,
        max_row=min(101, ws.max_row),
        max_col=min(ws.max_column, 200),
        values_only=True,
    ):
        rows.append(list(r))
    if not rows:
        return FileInspection(sheets, [], empty_long("col", 0), {}, warnings, {})
    header = ["" if v is None else str(v) for v in rows[0]]
    pre_merged = bool(ws.merged_cells.ranges) if hasattr(ws, "merged_cells") else False
    if pre_merged:
        warnings.append(
            ValidationIssue(
                code="E_XLSX_MERGED_CELL",
                severity=Severity.ERROR,
                message="Células mescladas encontradas na tabela; separe-as antes de importar.",
            )
        )
    for row in rows[1:6]:
        for col_idx, v in enumerate(row[: len(header)]):
            if isinstance(v, str) and v.startswith("="):
                warnings.append(
                    ValidationIssue(
                        code="E_XLSX_FORMULA_NOT_CACHED",
                        severity=Severity.ERROR,
                        message="Há fórmulas sem valor calculado; o Excel recalcula e "
                        "o app não executa fórmulas.",
                        field_name=header[col_idx] if col_idx < len(header) else None,
                    )
                )
    seen: dict[str, int] = {}
    dedup_header: list[str] = []
    for h in header:
        seen[h] = seen.get(h, 0) + 1
        dedup_header.append(h if seen[h] == 1 else f"{h}_{seen[h]}")
    sam = pl.DataFrame(
        [row[: len(header)] for row in rows[1:6]],
        schema={name: pl.Utf8 for name in dedup_header},
        orient="row",
    )
    types = {}
    return FileInspection(
        sheets,
        header,
        sam,
        types,
        warnings,
        {"n_rows_declared": ws.max_row, "n_cols": ws.max_column, "sheet": sheet},
    )
